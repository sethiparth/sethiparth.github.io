#!/usr/bin/env python3
"""Generate an ATS-safe one-page .docx resume for Parth Sethi.

ATS-safety principles (why this file is shaped the way it is):
- SINGLE COLUMN. No tables, no text boxes, no multi-column sections — ATS
  parsers read top-to-bottom and scramble anything else.
- Right-aligned dates via a RIGHT TAB STOP (not a table cell) so "Company ....
  Date" parses as one clean line.
- Standard section headings ("Experience", "Skills", ...) so parsers bucket
  content correctly.
- Contact details live in the BODY, never in the Word header/footer (many ATS
  ignore headers/footers entirely).
- Standard font (Calibri), real Unicode bullets, selectable text, no images.

Output: sethi_resume.docx at the repo root.
Run: .build/.venv/bin/python .build/build_resume.py
"""
import pathlib
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "sethi_resume.docx"

# --- palette / sizing -------------------------------------------------------
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # deep professional blue for headings
DIM = RGBColor(0x44, 0x44, 0x44)
FONT = "Calibri"

doc = Document()

# Tight one-page margins
sec = doc.sections[0]
sec.top_margin = Inches(0.5)
sec.bottom_margin = Inches(0.5)
sec.left_margin = Inches(0.6)
sec.right_margin = Inches(0.6)
CONTENT_WIDTH = sec.page_width - sec.left_margin - sec.right_margin

# Base style
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10)
normal.font.color.rgb = INK
pf = normal.paragraph_format
pf.space_before = Pt(0)
pf.space_after = Pt(0)
pf.line_spacing = 1.0


def set_spacing(p, before=0, after=0, line=1.0):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


import re as _re

def clean(text):
    """Map every ATS-risky non-ASCII glyph to a safe ASCII equivalent, so strict
    / legacy parsers (Taleo, older Workday configs) never mojibake or drop text.
    Smart quotes -> straight, em/en-dash -> hyphen, middot separator -> comma."""
    text = _re.sub(r"\s*·\s*", ", ", text)    # "  ·  " separator -> ", "
    text = _re.sub(r"\s*—\s*", " - ", text)   # em-dash (with surrounding spaces)
    return (text
            .replace("“", '"').replace("”", '"')   # curly double quotes
            .replace("‘", "'").replace("’", "'")   # curly single/apostrophe
            .replace("–", "-")                            # en-dash (date ranges)
            .replace("→", "to"))                          # arrow (belt-and-suspenders)


def run(p, text, size=10, bold=False, italic=False, color=INK, font=FONT):
    r = p.add_run(clean(text))
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    return r


def bottom_border(p):
    """Thin rule under a heading paragraph (via paragraph border XML)."""
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F4E79")
    pbdr.append(bottom)
    pPr.append(pbdr)


def section_heading(text):
    p = doc.add_paragraph()
    set_spacing(p, before=4, after=2)
    run(p, text.upper(), size=11, bold=True, color=ACCENT)
    bottom_border(p)
    return p


def two_col_line(left_runs, right_text, right_bold=False):
    """One paragraph: left-aligned content + right-aligned date via a right tab
    stop at the content edge (ATS-parseable, unlike a table)."""
    p = doc.add_paragraph()
    set_spacing(p, before=3, after=0)
    p.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    for text, kwargs in left_runs:
        run(p, text, **kwargs)
    run(p, "\t")
    run(p, right_text, size=9.5, bold=right_bold, color=DIM)
    return p


def bullet(text_segments):
    """A bullet line. text_segments = list of (text, kwargs) so we can bold
    metrics inline. Hanging indent so wrapped lines align under the text."""
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=1, line=1.0)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    run(p, "•  ", size=10)
    for text, kwargs in text_segments:
        run(p, text, **kwargs)
    return p


B = {"bold": True}
N = {}

# === HEADER =================================================================
name = doc.add_paragraph()
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(name, after=1)
run(name, "PARTH SETHI", size=22, bold=True, color=ACCENT)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(title, after=2)
run(title, "Senior Software Engineer", size=12, color=DIM)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(contact, after=2)
run(contact,
    "Seattle, WA  |  +1 (352) 745-9455  |  sethiparth.ps@gmail.com  |  "
    "https://linkedin.com/in/parth-sethi-a24753b7  |  https://sethiparth.github.io",
    size=9.5, color=DIM)

# === SUMMARY ================================================================
section_heading("Summary")
summ = doc.add_paragraph()
set_spacing(summ, before=2, after=1, line=1.05)
run(summ,
    "Senior Software Engineer, 6+ years designing high-throughput, fault-tolerant "
    "distributed systems and gRPC microservices at AWS and Salesforce. Operates at "
    "lead scope — architecting platforms end to end, coordinating cross-team "
    "launches, and setting the standards other teams build against. Recent work: "
    "the rendering engine serving a $1B+ revenue channel, a 10x throughput gain, "
    "and a 99.999% monthly availability SLO.", size=10)

# === EXPERIENCE =============================================================
section_heading("Professional Experience")

# --- Salesforce ---
two_col_line([("Salesforce", {"size": 10.5, "bold": True})],
             "May 2022 – Present", right_bold=True)
role = doc.add_paragraph(); set_spacing(role, after=1)
run(role, "Senior Software Engineer (SMTS)  ·  Seattle, WA", size=10, italic=True, color=DIM)

bullet([("Architected “Sage” from scratch", B),
        (", a two-step gRPC microservices rendering engine replacing the legacy "
         "E360 platform. Where the legacy library was a thin, low-level wrapper "
         "over Handlebars.js, Sage is a higher-order engine with deep, composable "
         "helpers; shipped two releases ahead of schedule.", N)])
bullet([("Led a cross-org effort across ", N), ("Q3 and Q4", B),
        (", coordinating ", N), ("4+ teams", B),
        (" and their dedicated test campaigns to clear Sage for launch — surfacing "
         "and fixing issues that hardened performance under peak request rates and "
         "heavy load.", N)])
bullet([("Led and orchestrated the migration of ", N), ("100% of customers", B),
        (" onto Sage with no downtime and no customer action required, reaching "
         "100% functional parity via a custom Compatibility Testing Framework "
         "built with Product.", N)])
bullet([("Set the reference architecture", B),
        (" now shared across ", N), ("3 channels", B),
        (" (SMS, email, RCS) as one engine — driving the design reviews and "
         "standards other teams build against.", N)])
bullet([("Built an agentic automation pipeline (Claude)", B),
        (" that replaced a manual ", N), ("~4-hour-per-org", B),
        (" setup — provisioning, test sends, and health checks — reclaiming "
         "engineer-weeks each quarter.", N)])
bullet([("Drove a ", N), ("10x throughput gain", B),
        (", 4M to 40M emails/hour, through DynamoDB batch tuning, Redis caching "
         "and deduplication, and quota-aware bulk sends.", N)])
bullet([("Built a pre-deploy validation framework", B),
        (" that caught regressions before rollout, sustaining a ", N), ("99.999%", B),
        (" monthly service-availability SLO; on call, cut incident response load "
         "and PagerDuty alerts by 10+/week.", N)])

# --- AWS ---
two_col_line([("Amazon Web Services (AWS)", {"size": 10.5, "bold": True})],
             "Oct 2019 – May 2022", right_bold=True)
role = doc.add_paragraph(); set_spacing(role, after=1)
run(role, "Software Development Engineer  ·  Seattle, WA", size=10, italic=True, color=DIM)

bullet([("Rearchitected the EC2 maintenance service", B),
        (" to fire “reboot” instance events across 25 public regions, protecting ", N),
        ("556,000+", B),
        (" customers’ instance data on degraded hardware.", N)])
bullet([("Redesigned EC2 maintenance scheduling", B),
        (" with a best-fit algorithm, eliminating daily manual scripts and "
         "cutting “no time slot” errors by 100s every two weeks.", N)])
bullet([("Automated manual maintenance escalations", B),
        (" on internal EC2 fleets, removing ", N), ("4,000 escalations/year", B),
        (" and hundreds of engineer-hours.", N)])
bullet([("Designed a customer-window scheduling algorithm", B),
        (" giving 10,000+ customers control over maintenance timing, plus an EC2 "
         "KPI “pain index” to predict and reduce escalations.", N)])

# --- OM Partners ---
two_col_line([("OM Partners", {"size": 10.5, "bold": True})],
             "Feb 2019 – Oct 2019", right_bold=True)
role = doc.add_paragraph(); set_spacing(role, after=1)
run(role, "Software Engineer  ·  Atlanta, GA", size=10, italic=True, color=DIM)

bullet([("Designed a task-distribution scheduling algorithm", B),
        (" over time buckets from user input, optimizing work-execution rate.", N)])
bullet([("Implemented optimization strategies in ", N), ("C++17", B),
        (" for “OPAL,” a custom supply-chain inventory and logistics language.", N)])

# === SKILLS =================================================================
section_heading("Technical Skills")
def skill_line(label, value):
    p = doc.add_paragraph(); set_spacing(p, before=1, after=1, line=1.0)
    run(p, label + ": ", size=10, bold=True)
    run(p, value, size=10)
skill_line("Languages", "Java, C++, SQL, NoSQL (proficient); C#, Scala, Python, Ruby (familiar)")
skill_line("AWS", "DynamoDB, Redis, Redshift, RDS, S3, Lambda, Kinesis, "
                  "QuickSight, EC2, CloudWatch, CloudFormation, IAM, KMS, SNS, SQS")
skill_line("Tools & Processes", "Git, GUS, Claude, Cursor, Agile/Kanban")

# === EDUCATION ==============================================================
section_heading("Education")
two_col_line([("M.S. Computer Science", {"size": 10, "bold": True}),
              ("  ·  University of Florida, Gainesville, FL", {"size": 10, "color": DIM})],
             "Aug 2017 – Dec 2018")
g = doc.add_paragraph(); set_spacing(g, after=1); run(g, "GPA 3.56 / 4.0", size=9.5, color=DIM)
two_col_line([("B.Tech Electronics & Communication", {"size": 10, "bold": True}),
              ("  ·  Jaypee University of Info. & Tech., Solan, India", {"size": 10, "color": DIM})],
             "Aug 2013 – May 2017")
g = doc.add_paragraph(); set_spacing(g, after=1); run(g, "GPA 9.1 / 10", size=9.5, color=DIM)

doc.save(str(OUT))
print(f"Wrote {OUT}")

# quick paragraph count as a sanity signal
print(f"{len(doc.paragraphs)} paragraphs")
